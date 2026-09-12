"""Title: Universal Converter

Tests for classify / extract and pipeline conversion of mixed file types.

Author: Eric Blaudez

Copyright (c) 2026 Thales
Licensed under the MIT License.
"""

from __future__ import annotations

import base64
import io
import zipfile

from thot.tasks.converters.Converter import Converter
from thot.tasks.converters.UniversalConverter import (
    UniversalConverter,
    _analyze_images_enabled,
    _captions_enabled,
    _decode_text,
    _embedded_image_limit,
    _extract_unknown,
    _legacy_doc_text,
    _mostly_printable,
    _ocr_heavy,
    _ocr_languages,
    _payload_is_analysable_raster,
    _payload_key,
    _pdf_images_per_page,
    _truncate,
    _xlsx_sheet_images,
    classify_kind,
    extract_markdown,
    pdf_images_analysis_markdown,
)


def _png_bytes(size: tuple[int, int] = (80, 80), color=(200, 20, 20)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def _with_zip_member(package: bytes, name: str, payload: bytes) -> bytes:
    source = io.BytesIO(package)
    dest = io.BytesIO()
    with (
        zipfile.ZipFile(source, "r") as zin,
        zipfile.ZipFile(dest, "w") as zout,
    ):
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(name, payload)
    return dest.getvalue()


class TestUniversalConverter:
    def test_classify_text_and_image_magic(self):
        assert classify_kind("note.txt", b"Hello") == "raw"
        assert classify_kind("scan.jpg", bytes([0xFF, 0xD8, 0xFF, 0xE0])) == (
            "image"
        )
        assert classify_kind("data.json", b'{"a": 1}') == "json"
        assert classify_kind("blob.bin", b"\x00\x01\xff") == "unknown"

    def test_extract_csv_markdown(self):
        text = extract_markdown(b"Name,Age\nAda,36\n", "t.csv", "csv")
        assert "Name" in text
        assert "Ada" in text

    def test_extract_html(self):
        html = b"<html><head><title>T</title></head><body>Hi</body></html>"
        text = extract_markdown(html, "p.html", "html")
        assert "Hi" in text

    def test_extract_zip_nested_text(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("inner.txt", "nested body")
        markdown = extract_markdown(buf.getvalue(), "pack.zip", "zip")
        assert "nested body" in markdown

    def test_convert_json_and_markdown(self):
        doc = UniversalConverter.convert(b'{"k": "v"}', "a.json", "json")
        assert "k" in doc["content"][0] or "v" in doc["content"][0]
        assert "k" in doc["markdown"]
        md = UniversalConverter.convert(b"# Title\n\nBody.\n", "a.md", "md")
        assert md["title"] == "Title"
        assert "Body." in md["content"][0]
        assert md["markdown"].startswith("# Title")

    def test_pipeline_auto_zip(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("readme.txt", "from archive")
        payload = base64.b64encode(buf.getvalue()).decode()
        document = Converter().convert(
            data_type="auto", data=payload, source="file://pack.zip"
        )
        joined = "\n".join(document["content"])
        assert "from archive" in joined
        assert document["conversion-info"]["datatype"] == "zip"

    def test_extract_image_markdown(self):
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (256, 256), (12, 80, 200)).save(buf, format="PNG")
        markdown = extract_markdown(
            buf.getvalue(),
            "chart.png",
            "image",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert markdown.startswith("# chart.png")
        assert "## Image analysis" in markdown
        assert "256x256" in markdown
        assert "scene" in markdown.lower()

    def test_pipeline_auto_image(self):
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (256, 256), "navy").save(buf, format="PNG")
        payload = base64.b64encode(buf.getvalue()).decode()
        document = Converter().convert(
            data_type="auto", data=payload, source="file://photo.png"
        )
        joined = "\n".join(document["content"])
        assert "Image analysis" in joined
        assert document["conversion-info"]["datatype"] == "image"

    def test_pipeline_auto_json(self):
        payload = base64.b64encode(b'{"hello": "world"}').decode()
        document = Converter().convert(
            data_type="auto", data=payload, source="file://rec.json"
        )
        joined = "\n".join(document["content"])
        assert "hello" in joined
        assert document["conversion-info"]["converter"] == "universal"

    def test_pdf_embedded_image_is_inline(self):
        import fitz
        from PIL import Image

        raster = io.BytesIO()
        Image.new("RGB", (320, 320), (20, 80, 180)).save(raster, format="PNG")
        document = fitz.open()
        page = document.new_page(width=400, height=600)
        page.insert_text((50, 40), "BEFORE FIGURE")
        page.insert_image(
            fitz.Rect(50, 80, 350, 380), stream=raster.getvalue()
        )
        page.insert_text((50, 430), "AFTER FIGURE")
        pdf = document.tobytes()
        document.close()
        markdown = extract_markdown(
            pdf,
            "figure.pdf",
            "pdf",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        before = markdown.find("BEFORE FIGURE")
        image = markdown.find("Embedded image")
        after = markdown.find("AFTER FIGURE")
        assert before != -1 and image != -1 and after != -1
        assert before < image < after

    def test_classify_kind_magic_and_suffixes(self):
        assert classify_kind("empty.txt", b"") == "raw"
        assert classify_kind("x.pdf", b"%PDF-1.4\n") == "pdf"
        assert classify_kind("x.rtf", b"{\\rtf1 hello}") == "rtf"
        assert classify_kind("x.html", b"<!DOCTYPE html><html></html>") == (
            "html"
        )
        assert classify_kind("plain.html", b"not markup") == "html"
        assert classify_kind("x.xml", b"<?xml version='1.0'?><a/>") == "xml"
        assert classify_kind("feed.rss", b"<?xml version='1.0'?><rss/>") == (
            "rss"
        )
        assert classify_kind("map.gml", b"<gml:Feature></gml:Feature>") == (
            "xml"
        )
        assert classify_kind("pic.tif", b"II*\x00") == "image"
        assert classify_kind("logo.svg", b"<svg></svg>") == "svg"
        assert classify_kind("n.docx", b"PK\x03\x04rest") == "docx"
        assert classify_kind("n.doc", b"xxxx") == "doc"
        assert classify_kind("n.pptx", b"PK\x03\x04rest") == "pptx"
        assert classify_kind("n.ppt", b"xxxx") == "ppt"
        assert classify_kind("n.xlsx", b"PK\x03\x04rest") == "xlsx"
        assert classify_kind("n.xls", b"xxxx") == "xls"
        assert classify_kind("n.tsv", b"a\tb\n") == "csv"
        assert classify_kind("note.md", b"# Hi") == "md"
        assert classify_kind("pack.zip", b"PK\x03\x04rest") == "zip"
        assert classify_kind("README", b"notes") == "raw"
        assert classify_kind("icon.bmp", b"BM") == "image"
        assert classify_kind("anim.gif", b"GIF89a....") == "image"
        assert classify_kind("pack.bin", b"PK\x03\x04rest") == "zip"
        assert classify_kind("plain", "notes".encode()) == "raw"
        assert classify_kind("odd.bin", b"\x80\x81\x82") == "unknown"
        assert _mostly_printable("ok")
        assert not _mostly_printable("")

    def test_extract_xml_svg_tsv_unknown_and_helpers(self):
        assert "hello" in extract_markdown(
            b"<root>hello</root>", "a.xml", "xml"
        )
        assert "circle" in extract_markdown(
            b"<svg><title>circle</title></svg>", "a.svg", "svg"
        )
        tsv = extract_markdown(b"Name\tAge\nAda\t36\n", "a.tsv", "tsv")
        assert "Ada" in tsv
        unknown = extract_markdown(b"\x00\x01ABC", "blob.bin", "unknown")
        assert "Unrecognized" in unknown
        assert "ABC" in _extract_unknown(b"ABC", "x.bin")
        assert _legacy_doc_text(b"") == ""
        assert _decode_text("café".encode()) == "café"
        assert "truncated at 3" in _truncate("abcd", 3)
        assert _payload_key(b"abc") == (3, b"abc")
        assert _ocr_languages(None).startswith("eng")
        assert _ocr_languages({"languages": "eng"}) == "eng"
        assert _analyze_images_enabled(None) is True
        assert _analyze_images_enabled({"analyze-images": False}) is False
        assert _ocr_heavy(None) is True
        assert _ocr_heavy({"enabled": False}) is False
        assert _captions_enabled(None) is True
        assert _captions_enabled({"enabled": True, "captions": False}) is False
        assert _captions_enabled({"enabled": True, "blip": False}) is False
        assert _embedded_image_limit(None) == 1024
        assert _embedded_image_limit({"max-embedded-images": 8}) == 8
        assert _pdf_images_per_page(None) == 32
        assert _pdf_images_per_page({"max-pdf-images-per-page": 4}) == 4
        assert "image" in UniversalConverter.managed_types()
        csv_rows = ["A,B"] + [f"{index},x" for index in range(502)]
        truncated = extract_markdown(
            "\n".join(csv_rows).encode(), "wide.csv", "csv"
        )
        assert "truncated" in truncated
        rtf = extract_markdown(b"{\\rtf1 hello}", "n.rtf", "rtf")
        assert "hello" in rtf
        assert "hello" in extract_markdown(b"hello", "a.txt", "auto")
        assert "hello" in extract_markdown(b"hello", "a.txt", "text")
        assert "Hi" in extract_markdown(
            b"<html><body>Hi</body></html>", "a.htm", "htm"
        )
        titled = extract_markdown(
            b"<html><head><title>UniqueTitleXYZ</title></head>"
            b"<body>Only body</body></html>",
            "t.html",
            "html",
        )
        assert "UniqueTitleXYZ" in titled
        assert extract_markdown(b"", "empty.csv", "csv") == ""
        ragged = extract_markdown(b"A,B\nC\n", "r.csv", "csv")
        assert "A" in ragged
        assert "item" in extract_markdown(
            b"<rss><item>item</item></rss>", "f.rss", "rss"
        )
        assert "node" in extract_markdown(
            b"<gml:Feature>node</gml:Feature>", "m.gml", "gml"
        )
        assert _decode_text(b"\x80") == "\x80"
        assert "Hi" in _decode_text("Hi".encode("utf-16"))
        assert isinstance(_legacy_doc_text(b"not a word document"), str)
        assert _payload_is_analysable_raster(b"not-image") is False

    def test_extract_docx_paragraph_and_picture(self):
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_paragraph("Hello Word body")
        picture = io.BytesIO(_png_bytes())
        document.add_picture(picture, width=Inches(1))
        document.add_paragraph("After figure")
        buf = io.BytesIO()
        document.save(buf)
        markdown = extract_markdown(
            buf.getvalue(),
            "note.docx",
            "docx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "Hello Word body" in markdown
        assert "After figure" in markdown

    def test_extract_docx_leftover_media(self):
        from docx import Document

        document = Document()
        document.add_paragraph("No inline pictures")
        buf = io.BytesIO()
        document.save(buf)
        packed = _with_zip_member(
            buf.getvalue(), "word/media/extra.png", _png_bytes()
        )
        markdown = extract_markdown(
            packed,
            "media.docx",
            "docx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "No inline pictures" in markdown or "Embedded image" in markdown

    def test_extract_pptx_slide_text_and_picture(self):
        from pptx import Presentation
        from pptx.util import Emu

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(
            Emu(0), Emu(0), Emu(1_000_000), Emu(400_000)
        )
        box.text_frame.text = "Hello Slide"
        raster = io.BytesIO(_png_bytes())
        slide.shapes.add_picture(
            raster, Emu(0), Emu(500_000), width=Emu(400_000)
        )
        buf = io.BytesIO()
        deck.save(buf)
        markdown = extract_markdown(
            buf.getvalue(),
            "talk.pptx",
            "pptx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "Hello Slide" in markdown
        assert "Slide 1" in markdown

    def test_extract_pptx_leftover_media(self):
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 1_000_000, 400_000)
        box.text_frame.text = "Deck notes"
        buf = io.BytesIO()
        deck.save(buf)
        packed = _with_zip_member(
            buf.getvalue(), "ppt/media/bg.png", _png_bytes()
        )
        markdown = extract_markdown(
            packed,
            "bg.pptx",
            "pptx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "Deck notes" in markdown or "Deck media" in markdown

    def test_extract_xlsx_table_empty_and_media(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage

        book = Workbook()
        people = book.active
        people.title = "People"
        people["A1"] = "Name"
        people["B1"] = "Age"
        people["A2"] = "Ada"
        people["B2"] = "36"
        empty = book.create_sheet("Empty")
        assert empty.title == "Empty"
        raster = io.BytesIO(_png_bytes())
        people.add_image(XLImage(raster), "C2")
        buf = io.BytesIO()
        book.save(buf)
        payload = buf.getvalue()
        markdown = extract_markdown(
            payload,
            "t.xlsx",
            "xlsx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "People" in markdown
        assert "Ada" in markdown
        assert "Empty" in markdown
        assert _xlsx_sheet_images(b"not-zip", "Sheet1", None) == ""
        packed = _with_zip_member(payload, "xl/media/chart.png", _png_bytes())
        media = extract_markdown(
            packed,
            "media.xlsx",
            "xlsx",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert "People" in media
        garbage = extract_markdown(b"not excel", "bad.xlsx", "xlsx")
        assert garbage == ""

    def test_convert_auto_empty_and_image_without_analysis(self):
        doc = UniversalConverter.convert(b"", "empty.txt", "auto")
        assert doc["source_doc_id"] == "empty.txt"
        markdown = extract_markdown(
            _png_bytes((256, 256), (0, 0, 80)),
            "icon.png",
            "image",
            ocr_config={"enabled": False, "analyze-images": False},
        )
        assert "icon.png" in markdown
        assert "256x256" in markdown
        tiny = extract_markdown(
            _png_bytes((32, 32)),
            "tiny.png",
            "image",
            ocr_config={"enabled": False, "analyze-images": True},
        )
        assert tiny == ""

    def test_bad_zip_and_macos_members(self):
        markdown = extract_markdown(b"PK\x03\x04notazip", "x.zip", "zip")
        assert "Unrecognized" in markdown
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as archive:
            archive.writestr("keep.txt", "visible")
            archive.writestr("__MACOSX/._keep.txt", "hidden")
            archive.writestr("._skip.txt", "appledouble")
            archive.writestr("nested/", "")
        zipped = extract_markdown(buf.getvalue(), "pack.zip", "zip")
        assert "visible" in zipped
        assert "appledouble" not in zipped

    def test_extract_doc_kind_and_xls_fallback(self):
        from docx import Document
        from openpyxl import Workbook

        document = Document()
        document.add_paragraph("Legacy path")
        buf = io.BytesIO()
        document.save(buf)
        markdown = extract_markdown(buf.getvalue(), "note.doc", "doc")
        assert isinstance(markdown, str)
        book = Workbook()
        book.active["A1"] = "X"
        xlsx = io.BytesIO()
        book.save(xlsx)
        fallback = extract_markdown(xlsx.getvalue(), "t.xls", "xls")
        assert isinstance(fallback, str)

    def test_pdf_scanned_page_and_public_wrapper(self):
        import fitz

        document = fitz.open()
        page = document.new_page(width=200, height=200)
        page.insert_text((20, 40), "x")
        pdf = document.tobytes()
        document.close()
        markdown = extract_markdown(
            pdf,
            "scan.pdf",
            "pdf",
            ocr_config={
                "enabled": False,
                "analyze-images": True,
                "min-page-text-chars": 400,
                "render-dpi": 72,
            },
        )
        assert "Full-page analysis" in markdown or "x" in markdown
        assert (
            pdf_images_analysis_markdown(pdf, {"analyze-images": False}) == ""
        )
        assert extract_markdown(b"%PDF-1.4 junk", "bad.pdf", "pdf") == ""
